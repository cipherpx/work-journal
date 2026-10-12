from core.cfg_manager import get_root_path, load_config_toml
from core.log_service import query_by_day
from core.gen_report import gen_daily_report, gen_weekly_report
from core.utils import make_dirs, now_ts_file_str
import os
import pyperclip
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = get_root_path()
EXPORT_DIR = os.path.join(ROOT, str(load_config_toml()['path']['export_dir']))
make_dirs(EXPORT_DIR)

def export_to_txt(content, filename=None):
    if not filename:
        filename = f"report_{now_ts_file_str()}.txt"
    path = os.path.join(EXPORT_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    return path

def export_to_excel(year, month, day):
    items = query_by_day(year, month, day)
    wb = Workbook()
    ws = wb.active
    ws.title = "日志明细"
    headers = ["ID", "日期", "标题", "内容", "优先级", "重要度", "星级", "分类"]
    ws.append(headers)
    for it in items:
        ws.append([
            it[0], f"{it[2]}-{it[3]:02d}-{it[4]:02d}",
            it[6], it[7], it[8], it[9], it[10], it[11]
        ])
    fn = f"log_{year}{month:02d}{day:02d}.xlsx"
    path = os.path.join(EXPORT_DIR, fn)
    wb.save(path)
    return path

def export_to_word(content, filename=None):
    if not filename:
        filename = f"report_{now_ts_file_str()}.docx"
    path = os.path.join(EXPORT_DIR, filename)
    doc = Document()
    for para in content.split("\n"):
        doc.add_paragraph(para)
    doc.save(path)
    return path

def export_to_pdf(content, filename=None):
    if not filename:
        filename = f"report_{now_ts_file_str()}.pdf"
    path = os.path.join(EXPORT_DIR, filename)
    # 注册中文字体（确保系统有该字体）
    try:
        pdfmetrics.registerFont(TTFont("SimSun", "C:/Windows/Fonts/simsun.ttc"))
    except:
        pass
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("SimSun", 10)
    y = 800
    for line in content.split("\n"):
        c.drawString(50, y, line)
        y -= 20
        if y < 50:
            c.showPage()
            c.setFont("SimSun", 10)
            y = 800
    c.save()
    return path

def copy_to_clipboard(content):
    pyperclip.copy(content)
    return True

# 对外快捷方法
def export_daily_all(year, month, day):
    content = gen_daily_report(year, month, day)
    if len(content.split('\n')) <= 6:
        raise ValueError('当日不存在可用日志！')
    return {
        "txt": export_to_txt(content),
        "xlsx": export_to_excel(year, month, day),
        "docx": export_to_word(content),
        "pdf": export_to_pdf(content),
        "clipboard": copy_to_clipboard('\n'.join(content.split('\n')[:-5]))
    }

def export_weekly_all(year, week):
    content = gen_weekly_report(year, week)
    if len(content.split('\n')) <= 5:
        raise ValueError('当周不存在可用日志！')
    return {
        "txt": export_to_txt(content),
        "docx": export_to_word(content),
        "pdf": export_to_pdf(content),
        "clipboard": copy_to_clipboard(content)
    }