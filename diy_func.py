#! python3
# diy_func.py
# 实现功能：根据自身需求，补充的功能模块
# Author：cipherpx
# Date: 2026/5/15 09:26
import shutil

import re
import tkinter as tk
import os
import subprocess
import platform
import toml
import sys
from PyQt5.QtWidgets import QApplication, QWidget, QDateEdit, QPushButton, QVBoxLayout
from PyQt5.QtCore import QDate
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from pathlib import Path

#由绝对路径修改为相对路径
TOML_PATH = Path(__file__).parent / 'config.toml'
TOML_EXAMPLE = Path(__file__).parent / 'config.toml.example'

def update_config_category(new_options):
    # 加载原有配置
    cfg = load_config_toml()
    # 去重、保留顺序，强制保留末尾 '自定义'
    unique_list = []
    for item in new_options:
        item = item.strip()
        if item and item not in unique_list and item != '自定义':
            unique_list.append(item)
    # 末尾固定加自定义
    unique_list.append('自定义')
    # 写回配置
    cfg['category']['options'] = unique_list
    update_config_toml(cfg)
    return unique_list

def add_category_if_not_exists(new_category):
    # 🔥 先判断 None，绝对安全
    if new_category is None:
        return
    # 去除空格
    new_category = new_category.strip()
    # 空字符串也不处理
    if not new_category:
        return

    # 加载你的配置
    config = load_config_toml()

    if new_category not in config['category']['options']:
        config['category']['options'].insert(-1, new_category)
        # 保存
    update_config_toml(config)

def load_config_toml(config_path=TOML_PATH,config_example=TOML_EXAMPLE):
    if not os.path.exists(config_path):
        config_dir = os.path.dirname(config_path)
        if config_dir and not os.path.exists(config_dir):
            os.makedirs(config_dir,exist_ok=True)
        if os.path.exists(config_example):
            shutil.copy(config_example,config_path)
        else:
            raise FileNotFoundError( f'配置文件 {config_path} 不存在，且示例文件 {config_example} 也不存在，无法生成默认配置。')
    return toml.load(config_path)

def replace_docx_content(doc, replace_map: dict):
    # 过滤值为None的占位，不处理
    valid_map = {k: str(v) for k, v in replace_map.items() if v is not None and v != ''}
    if not valid_map:
        return

    def insert_paragraph_after(target_p: Paragraph, text: str = '') -> Paragraph:
        """在指定段落后面插入全新独立段落（硬回车）"""
        # 创建新段落XML节点，插在目标段落后
        new_p_xml = OxmlElement("w:p")
        target_p._element.addnext(new_p_xml)
        new_p = Paragraph(new_p_xml, target_p._parent)

        # 复制原段落整体格式（样式、缩进、行距、段间距）
        new_p.style = target_p.style
        new_p.paragraph_format.alignment = target_p.paragraph_format.alignment
        new_p.paragraph_format.left_indent = target_p.paragraph_format.left_indent
        new_p.paragraph_format.right_indent = target_p.paragraph_format.right_indent
        new_p.paragraph_format.first_line_indent = target_p.paragraph_format.first_line_indent
        new_p.paragraph_format.space_before = target_p.paragraph_format.space_before
        new_p.paragraph_format.space_after = target_p.paragraph_format.space_after
        new_p.paragraph_format.line_spacing = target_p.paragraph_format.line_spacing

        # 写入文字并复制字体
        if text:
            run = new_p.add_run(text)
            # 复制第一段文字字体格式
            if target_p.runs:
                src_run = target_p.runs[0]
                run.font.name = src_run.font.name
                run.font.size = src_run.font.size
                run.font.bold = src_run.font.bold
                run.font.italic = src_run.font.italic
                run.font.underline = src_run.font.underline
                run.font.color.rgb = src_run.font.color.rgb
                # 底层同步中英西文字体
                src_rpr = src_run._element.get_or_add_rPr()
                src_rfonts = src_rpr.get_or_add_rFonts()
                dst_rpr = run._element.get_or_add_rPr()
                dst_rfonts = dst_rpr.get_or_add_rFonts()
                # 复制各类字体
                ascii_val = src_rfonts.get(qn("w:ascii"))
                if ascii_val:
                    run.font.name = ascii_val
                # 复制中文字体
                ea_val = src_rfonts.get(qn("w:eastAsia"))
                if ea_val:
                    dst_rfonts.set(qn("w:eastAsia"), ea_val)
        return new_p

    def replace_paragraph(para):
        # 合并当前段落所有文字
        full_text = ''.join(run.text for run in para.runs)
        # 批量替换
        for old, new in valid_map.items():
            full_text = full_text.replace(old, new)
        # 拆分换行文本，生成多个段落
        text_lines = full_text.split('\n')
        # 清空原有run，写入新文本（保留段落格式）
        for run in para.runs:
            run.text = ''
        if para.runs:
            para.runs[0].text = text_lines[0]
        else:
            para.add_run(text_lines[0])
        # 剩余每行新建段落
        if len(text_lines) > 1:
            new_p = para
            for line in text_lines[1:]:
                new_p = insert_paragraph_after(new_p,line)

    # 1. 正文段落
    for para in doc.paragraphs:
        replace_paragraph(para)

    # 2. 表格单元格段落
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_paragraph(para)


def update_config_toml(cfg,config_path=TOML_PATH):
    with open(config_path, 'w', encoding='utf-8') as f:
        toml.dump(cfg, f)



# ==============================================
# 独立工具函数：调用就弹出日历选择器
# ==============================================
def pick_date():
    """
    弹出日历选择器，返回选中的日期字符串（yyyy-MM-dd）
    用法：date = pick_date()
    """
    app = QApplication.instance()
    if not app:
        app = QApplication(sys.argv)

    result = None

    # 弹窗窗口
    win = QWidget()
    win.setWindowTitle("选择日期")
    win.resize(360, 180)
    win.setStyleSheet("background-color:#f8f9fa;")

    # 布局
    layout = QVBoxLayout(win)
    layout.setContentsMargins(30, 25, 30, 25)
    layout.setSpacing(15)

    # 日期选择框（带日历弹窗）
    date_edit = QDateEdit()
    date_edit.setDate(QDate.currentDate())
    date_edit.setCalendarPopup(True)
    date_edit.setDisplayFormat("yyyy-MM-dd")
    date_edit.setStyleSheet("""
        QDateEdit {
            height:34px;
            font-size:14px;
            padding-left:10px;
            border:1px solid #ced4da;
            border-radius:6px;
        }
    """)
    layout.addWidget(date_edit)

    # 确定按钮
    btn = QPushButton("确认选择")
    btn.setStyleSheet("""
        QPushButton {
            height:36px;
            background-color:#409EFF;
            color:white;
            border:none;
            border-radius:6px;
            font-size:14px;
        }
        QPushButton:hover {
            background-color:#66b1ff;
        }
    """)
    layout.addWidget(btn)

    # 按钮点击逻辑
    def on_ok():
        nonlocal result
        result = date_edit.date().toString("yyyy-MM-dd")
        win.close()

    btn.clicked.connect(on_ok)
    win.show()
    app.exec_()

    return result

def parse_and_fix_date(date_str):
    """
    智能解析日期：
    - 轻微不合法：自动修复（如 2025-01-02 多空格、横杠错误、少零）
    - 严重不合法：返回 None
    """
    date_str = str(date_str).strip()

    # 1. 把所有分隔符统一换成 -
    date_str = re.sub(r'[./\\]', '-', date_str)
    # 去掉多余空格
    date_str = re.sub(r'\s+', '', date_str)

    # 匹配 YYYY-M-D 或 YYYY-MM-DD
    match = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})$', date_str)
    if not match:
        return None

    y, m, d = match.groups()

    try:
        y = int(y)
        m = int(m)
        d = int(d)
    except:
        return None

    # 简单范围校验
    if not (2000 <= y <= 2100 and 1 <= m <= 12 and 1 <= d <= 31):
        return None

    # 自动补零，返回标准格式：YYYY-MM-DD
    return f"{y:04d}-{m:02d}-{d:02d}"

# 跨平台打开文件夹
def open_folder(path):
    if not os.path.exists(path):
        return
    if platform.system() == "Windows":
        os.startfile(path)
    elif platform.system() == "Darwin":
        subprocess.run(["open", path])
    else:
        subprocess.run(["xdg-open", path])

# 跨平台直接打开文件
def open_file(path):
    if not os.path.exists(path):
        return
    if platform.system() == "Windows":
        os.startfile(path)
    elif platform.system() == "Darwin":
        subprocess.run(["open", path])
    else:
        subprocess.run(["xdg-open", path])

# 封装保存成功弹窗
def show_save_success(parent: tk.Tk, file_path: str,text='通知文本',title='保存成功'):
    dir_path = os.path.dirname(file_path)
    win = tk.Toplevel(parent)
    win.title(title)
    win.geometry("520x180")
    win.focus_force()
    win.grab_set()

    # ========== 弹窗居中主窗口核心代码 ==========
    parent.update()
    # 获取主窗口尺寸、坐标
    main_w = parent.winfo_width()
    main_h = parent.winfo_height()
    main_x = parent.winfo_x()
    main_y = parent.winfo_y()
    # 弹窗宽高
    pop_w = 520
    pop_h = 180
    # 计算居中坐标
    x = main_x + (main_w - pop_w) // 2
    y = main_y + (main_h - pop_h) // 2
    win.geometry(f"{pop_w}x{pop_h}+{x}+{y}")

    # ==========================================

    # 统一关闭函数，彻底销毁无残留
    def close_all():
        win.destroy()

    win.protocol("WM_DELETE_WINDOW", close_all)

    tk.Label(win, text=text, wraplength=490, font=("", 10)).pack(pady=22)
    btn_frame = tk.Frame(win)
    btn_frame.pack()

    tk.Button(btn_frame, text="确定", width=10, command=close_all).grid(row=0, column=0, padx=7)
    tk.Button(btn_frame, text="打开文件夹", width=10,
              command=lambda: (open_folder(dir_path), close_all())).grid(row=0, column=1, padx=7)
    tk.Button(btn_frame, text="打开文件", width=10,
              command=lambda: (open_file(file_path), close_all())).grid(row=0, column=2, padx=7)

    win.wait_window(win)


# ==============================================
# 主函数直接调用！！！
# ==============================================
if __name__ == '__main__':
    # 一行调用日历
    selected_date = pick_date()

    # 输出结果
    print("你选中的日期是：", selected_date)